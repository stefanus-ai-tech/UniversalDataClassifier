from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterator

SUPPORTED = {".csv", ".tsv", ".json", ".jsonl", ".txt", ".log", ".md", ".pgn", ".xlsx"}


def _looks_like_pgn(path: Path) -> bool:
    if path.suffix.lower() not in {".txt", ".pgn"}:
        return False
    with path.open("r", encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(12000)
    return bool(re.search(r'^\[(?:Event|Site|GameID)\s+"[^"]*"\]', sample, re.M) and re.search(r'^\[Result\s+"(?:1-0|0-1|1/2-1/2|\*)"\]', sample, re.M))


def _pgn_records(path: Path) -> Iterator[tuple[int, dict]]:
    tags: dict[str, str] = {}
    moves: list[str] = []
    index = 0
    for line in _text_lines(path):
        stripped = line.strip()
        match = re.fullmatch(r'\[([A-Za-z][A-Za-z0-9_]*)\s+"(.*)"\]', stripped)
        if match:
            if moves:
                index += 1
                yield index, {**tags, "moves": " ".join(moves)}
                tags, moves = {}, []
            tags[match.group(1)] = match.group(2)
        elif stripped and tags:
            moves.append(stripped)
    if tags:
        index += 1
        yield index, {**tags, "moves": " ".join(moves)}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return "sha256:" + digest.hexdigest()


def _text_lines(path: Path):
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            yield from handle
    except UnicodeDecodeError as exc:
        raise ValueError("File text must be UTF-8 encoded") from exc


def _json_row(value):
    if isinstance(value, dict):
        return value
    return {"value": value}


def read_records(path: Path) -> Iterator[tuple[int, dict]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported format: {suffix}")
    if _looks_like_pgn(path):
        yield from _pgn_records(path)
    elif suffix in {".csv", ".tsv"}:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if suffix == ".csv" else csv.excel_tab
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(handle, dialect=dialect)
            if not reader.fieldnames:
                return
            for row_number, row in enumerate(reader, 2):
                yield row_number, {str(k): v for k, v in row.items() if k is not None}
    elif suffix == ".jsonl":
        for row_number, line in enumerate(_text_lines(path), 1):
            if line.strip():
                try:
                    yield row_number, _json_row(json.loads(line))
                except json.JSONDecodeError as exc:
                    yield row_number, {"_parse_error": str(exc), "raw_line": line.strip()}
    elif suffix == ".json":
        if path.stat().st_size > 20 * 1024 * 1024:
            raise ValueError("JSON arrays above 20 MB should be converted to JSONL for streaming")
        with path.open("r", encoding="utf-8-sig") as handle:
            content = json.load(handle)
        if isinstance(content, dict):
            content = content.get("records", [content])
        if not isinstance(content, list):
            content = [content]
        for row_number, value in enumerate(content, 1):
            yield row_number, _json_row(value)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        ordinal = 0
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                names = [str(cell) if cell is not None else f"column_{i+1}" for i, cell in enumerate(header)]
                for row_number, values in enumerate(rows, 2):
                    ordinal += 1
                    yield ordinal, {"_source_sheet": sheet.title, "_source_row": row_number, **{names[i]: value for i, value in enumerate(values) if value is not None}}
        finally:
            workbook.close()
    else:
        for row_number, line in enumerate(_text_lines(path), 1):
            text = line.strip()
            if text:
                yield row_number, {"text": text}


def profile(path: Path, sample_limit: int = 12) -> dict:
    fields: dict[str, dict] = defaultdict(lambda: {"present": 0, "numeric": 0, "examples": Counter(), "length_sum": 0})
    count = 0
    duplicates = 0
    seen: set[str] = set()
    samples = []
    malformed = 0
    for _, row in read_records(path):
        count += 1
        if "_parse_error" in row:
            malformed += 1
        if len(samples) < sample_limit:
            samples.append(row)
        fingerprint = hashlib.sha256(json.dumps(row, sort_keys=True, ensure_ascii=False, default=str).encode()).digest()
        if fingerprint in seen:
            duplicates += 1
        elif len(seen) < 100_000:
            seen.add(fingerprint)
        for key, value in row.items():
            if value is None or str(value).strip() == "":
                continue
            field = fields[key]
            field["present"] += 1
            field["length_sum"] += len(str(value))
            if len(field["examples"]) < 100 or str(value) in field["examples"]:
                field["examples"][str(value)[:100]] += 1
            if isinstance(value, (int, float)) or re.fullmatch(r"[+-]?\d+(?:\.\d+)?", str(value).strip()):
                field["numeric"] += 1
    columns = []
    for key, stats in fields.items():
        present = stats["present"]
        columns.append({
            "name": key,
            "type": "number" if present and stats["numeric"] / present > .9 else "text",
            "missing_pct": round((1 - present / count) * 100, 1) if count else 0,
            "unique_sample": len(stats["examples"]),
            "avg_length": round(stats["length_sum"] / present, 1) if present else 0,
            "top_values": stats["examples"].most_common(3),
        })
    return {"filename": path.name, "format": "pgn" if _looks_like_pgn(path) else path.suffix.lower().lstrip("."), "source_hash": sha256(path), "records": count, "field_count": len(columns), "duplicate_count": duplicates, "malformed_count": malformed, "columns": columns, "samples": samples}
