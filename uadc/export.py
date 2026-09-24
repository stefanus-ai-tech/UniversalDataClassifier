from __future__ import annotations

import json
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .pipeline import RUNS, get_run, summary

NAVY = "17233A"
TEAL = "19B8A7"
PALE = "EAF7F5"
WHITE = "FFFFFF"


def _sheet(book, name, headers):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    return sheet


def _finish(sheet):
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        max_len = max((len(str(cell.value or "")) for cell in list(column)[:200]), default=8)
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 55)
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def make_workbook(run_id: str) -> Path:
    run = get_run(run_id)
    if run["status"] != "COMPLETE":
        raise ValueError("Export is available after the run completes")
    stats = summary(run_id)
    book = Workbook()
    dashboard = book.active
    dashboard.title = "01_Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:F2")
    dashboard["A1"] = "UNIVERSAL ADAPTIVE DATA CLASSIFIER"
    dashboard["A1"].font = Font(size=18, bold=True, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard["A3"] = f"Run {run_id}"
    metrics = [
        ("Input records", run["metrics"]["input"]),
        ("Processed", run["metrics"]["processed"]),
        ("Auto approved", stats["counts"].get("AUTO_APPROVED", 0)),
        ("Needs review", stats["counts"].get("NEEDS_REVIEW", 0)),
        ("Unresolved", stats["counts"].get("UNRESOLVED", 0)),
        ("Duplicates removed", run["metrics"]["duplicates"]),
        ("Rejected", run["metrics"]["rejected"]),
        ("Avg model confidence", stats["average_model_confidence"]),
    ]
    for index, (label, value) in enumerate(metrics, 5):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, value)
        dashboard.cell(index, 1).font = Font(bold=True, color=NAVY)
        dashboard.cell(index, 2).fill = PatternFill("solid", fgColor=PALE)
    dashboard["D5"] = "Label"
    dashboard["E5"] = "Records"
    for index, (label, count) in enumerate(stats["labels"].items(), 6):
        dashboard.cell(index, 4, label)
        dashboard.cell(index, 5, count)
    if stats["labels"]:
        chart = BarChart()
        chart.title = "Classification distribution"
        chart.add_data(Reference(dashboard, min_col=5, min_row=5, max_row=5 + len(stats["labels"])), titles_from_data=True)
        chart.set_categories(Reference(dashboard, min_col=4, min_row=6, max_row=5 + len(stats["labels"])))
        dashboard.add_chart(chart, "D14")
    dashboard.column_dimensions["A"].width = 27
    dashboard.column_dimensions["B"].width = 18
    dashboard.column_dimensions["D"].width = 25
    dashboard.column_dimensions["E"].width = 16

    classifications = _sheet(book, "02_Classifications", ["Record ID", "Source row", "Label", "Confidence", "Engine", "Review", "Action", "Alternatives"])
    quality = _sheet(book, "03_Data_Quality", ["Field", "Type", "Missing %", "Unique sample", "Avg length", "Top values"])
    confidence = _sheet(book, "04_Confidence", ["Record ID", "Label", "Model confidence", "Review status"])
    cleaning = _sheet(book, "05_Cleaning_Audit", ["Record ID", "Field", "Before", "Operation", "After"])
    semantic = _sheet(book, "06_Semantic_State", ["Record ID", "Fact", "Value", "Evidence"])
    actions = _sheet(book, "07_API_Actions", ["Record ID", "Label", "Method", "Endpoint", "Request", "Response code", "Response", "Status"])
    review_queue = _sheet(book, "08_Review_Queue", ["Record ID", "Prediction", "Confidence", "Reason", "Raw preview", "Cleaned preview"])
    errors = _sheet(book, "09_Errors", ["Record ID", "Stage", "Error"])
    dataset = _sheet(book, "10_Dataset_Profile", ["Property", "Value"])
    metadata = _sheet(book, "11_Run_Metadata", ["Property", "Value"])
    for column in run["profile"]["columns"]:
        quality.append([column["name"], column["type"], column["missing_pct"], column["unique_sample"], column["avg_length"], json.dumps(column["top_values"], ensure_ascii=False)])
    for key in ("filename", "format", "source_hash", "records", "field_count", "duplicate_count", "malformed_count"):
        dataset.append([key, run["profile"][key]])
    for key in ("id", "created_at", "completed_at", "contract", "contract_id", "objective", "status"):
        metadata.append([key, run.get(key)])
    metadata.append(["cleaning_plan", json.dumps(run["plan"], ensure_ascii=False)])
    connection = sqlite3.connect(RUNS / f"{run_id}.sqlite")
    try:
        for row in connection.execute("SELECT * FROM records ORDER BY row_number"):
            record_id, row_number = row[0], row[1]
            raw, cleaned, facts, evidence, decision, review, action, audit = [json.loads(value) for value in row[2:]]
            label, score = decision.get("label"), decision.get("confidence")
            classifications.append([record_id, row_number, label, score, decision["engine"], review["status"], action["status"], json.dumps(decision.get("alternatives", {}), ensure_ascii=False)])
            confidence.append([record_id, label, score, review["status"]])
            for change in audit:
                cleaning.append([record_id, change["field"], str(change["before"]), change["operation"], str(change["after"])])
            for fact, value in facts.items():
                semantic.append([record_id, fact, str(value)[:1000], json.dumps(evidence.get(fact, []), ensure_ascii=False)[:2000]])
            actions.append([record_id, label, action.get("method"), action.get("endpoint"), json.dumps(action.get("request"), ensure_ascii=False), action.get("response_code"), json.dumps(action.get("response"), ensure_ascii=False), action["status"]])
            if review["status"] != "AUTO_APPROVED":
                review_queue.append([record_id, label, score, review["reason"], json.dumps(raw, ensure_ascii=False)[:500], json.dumps(cleaned, ensure_ascii=False)[:500]])
            if decision.get("error"):
                errors.append([record_id, "classification", decision["error"]])
    finally:
        connection.close()
    for sheet in book.worksheets[1:]:
        _finish(sheet)
    output = RUNS / f"{run_id}.xlsx"
    book.save(output)
    return output
