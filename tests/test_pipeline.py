import io
import json
import os
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import app
from uadc.cleaning import clean_record, validate_plan
from uadc.contracts import load_contract
from uadc.decision import gate
from uadc.decision import classify
from uadc.ingest import profile, read_records
from uadc.pipeline import contract_warning
from pathlib import Path
from tempfile import TemporaryDirectory


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_csv_to_excel_and_review_queue(self):
        data = "ticket_id,message,phone\nT1,  payment problem  ,0812 3456 7890\nT2,login error,08122223333\nT2,login error,08122223333\n"
        with patch.dict(os.environ, {"GROQ_API_KEY": "", "LAYA_URL": ""}):
            intake = self.client.post("/api/intake", files={"file": ("tickets.csv", data.encode(), "text/csv")}, data={"contract": "support"})
            self.assertEqual(intake.status_code, 200, intake.text)
            run = intake.json()
            self.assertEqual(run["profile"]["records"], 3)
            self.assertEqual(run["status"], "READY")
            execute = self.client.post(f"/api/runs/{run['id']}/execute")
            self.assertEqual(execute.status_code, 200, execute.text)
            detail = self.client.get(f"/api/runs/{run['id']}").json()
            self.assertEqual(detail["status"], "COMPLETE", detail.get("error"))
            self.assertEqual(detail["metrics"]["duplicates"], 1)
            records = self.client.get(f"/api/runs/{run['id']}/records").json()["records"]
            self.assertEqual(len(records), 2)
            self.assertEqual(records[0]["cleaned_data"]["phone"], "+6281234567890")
            self.assertEqual(records[0]["review"]["status"], "UNRESOLVED")
            self.assertEqual(records[0]["action"]["status"], "NOT_SENT")
            response = self.client.get(f"/api/runs/{run['id']}/export/xlsx")
            self.assertEqual(response.status_code, 200, response.text[:500] if response.status_code != 200 else "")
            book = load_workbook(io.BytesIO(response.content), read_only=True)
            self.assertEqual(len(book.sheetnames), 11)
            self.assertEqual(book["02_Classifications"].max_row, 3)
            self.assertEqual(book["08_Review_Queue"].max_row, 3)
            csv_response = self.client.get(f"/api/runs/{run['id']}/export/csv")
            self.assertEqual(csv_response.status_code, 200)
            self.assertEqual(len(csv_response.text.splitlines()), 3)
            json_response = self.client.get(f"/api/runs/{run['id']}/export/json")
            self.assertEqual(len(json_response.json()["records"]), 2)

    def test_plan_rejects_untrusted_code(self):
        with self.assertRaises(ValueError):
            validate_plan({"operations": [{"field": "message", "op": "eval"}]}, {"message"})
        cleaned, audit = clean_record({"message": "  hello   world "}, {"operations": [{"field": "message", "op": "trim_whitespace"}]})
        self.assertEqual(cleaned["message"], "hello world")
        self.assertEqual(len(audit), 1)

    def test_confidence_gate(self):
        contract = load_contract("support")
        self.assertEqual(gate({"engine": "laya", "label": "billing", "confidence": .9}, contract)["status"], "AUTO_APPROVED")
        self.assertEqual(gate({"engine": "laya", "label": "billing", "confidence": .7}, contract)["status"], "NEEDS_REVIEW")
        self.assertEqual(gate({"engine": "rules_demo", "label": "billing", "confidence": None}, contract)["status"], "UNRESOLVED")

    def test_laya_response_contract(self):
        class FakeResponse:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self, *args):
                return json.dumps({"answers": {"department": {"choice": "billing", "confidence": .7, "answer_confidence": .93, "probabilities": {"billing": .93, "account": .07}}, "urgency": {"score": 2.0}}, "routing": {"model": "multilingual"}}).encode()

        with patch.dict(os.environ, {"LAYA_URL": "http://127.0.0.1:8001"}), patch("urllib.request.urlopen", return_value=FakeResponse()) as mocked:
            result = classify({"facts": {"message": "charged twice"}}, load_contract("support"))
        self.assertEqual(result["label"], "billing")
        self.assertEqual(result["confidence"], .93)
        self.assertEqual(result["routing"]["model"], "multilingual")
        self.assertEqual(gate(result, load_contract("support"))["status"], "AUTO_APPROVED")
        self.assertIn("/v1/systemone", mocked.call_args.args[0].full_url)

    def test_pgn_is_grouped_by_game_and_mismatch_is_flagged(self):
        data = '[Site "VRChess"]\n[Result "1-0"]\n\n1.e4 e5 1-0\n\n[Site "VRChess"]\n[Result "0-1"]\n\n1.d4 d5 0-1\n'
        with TemporaryDirectory() as folder:
            path = Path(folder) / "games.txt"
            path.write_text(data, encoding="utf-8")
            records = list(read_records(path))
            self.assertEqual(len(records), 2)
            self.assertEqual(records[0][1]["Result"], "1-0")
            data_profile = profile(path)
            self.assertEqual(data_profile["format"], "pgn")
            self.assertIsNotNone(contract_warning(data_profile, "support"))

    def test_offline_laya_is_reported_before_processing(self):
        with patch.dict(os.environ, {"GROQ_API_KEY": "", "LAYA_URL": ""}):
            created = self.client.post("/api/intake", files={"file": ("tickets.csv", b"message\nrefund please\n", "text/csv")}, data={"contract": "support"}).json()
        with patch("app.laya_health", return_value={"configured": True, "reachable": False, "message": "Laya server is unreachable"}):
            response = self.client.post(f"/api/runs/{created['id']}/execute")
        self.assertEqual(response.status_code, 503)
        self.assertIn("Laya server is unreachable", response.json()["detail"])
        self.assertEqual(self.client.get(f"/api/runs/{created['id']}").json()["status"], "READY")


if __name__ == "__main__":
    unittest.main()
